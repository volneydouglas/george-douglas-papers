from pathlib import Path
import json, re, csv, hashlib, shutil, unicodedata, collections
from docx import Document
from pypdf import PdfReader
import pdfplumber
from lxml import html

BASE=Path('/Users/volneydouglas/Documents/FamilyHistory')
WORK=BASE/'tmp/george_douglas/revision'
OUT=BASE/'output/george-douglas-archival'
SUP=OUT/'Supporting-files'
obj=json.loads((SUP/'Transcripts.json').read_text())
pages=obj['pages']; sources=obj['sources']
expected=[f"{s['id']}-{n:03}" for s in sources for n in range(1,s['pages']+1)]
assert [p['id'] for p in pages]==expected
assert len(pages)==387
source_root=Path('/Users/volneydouglas/Library/Mobile Documents/com~apple~CloudDocs/George Douglas')
for source in sources:
    source_file=source_root/source['filename']
    assert hashlib.sha256(source_file.read_bytes()).hexdigest()==source['sha256']
    assert len(PdfReader(source_file).pages)==source['pages']

def norm(s):
    return re.sub(r'\s+', '',unicodedata.normalize('NFKC',s)).replace('\u00ad','')

page_map=[]; volume_counts={}; pdf_issues=[]
block_ids=[b['id'] for p in pages for b in p['paragraphs']]
assert len(block_ids)==len(set(block_ids))==971
for roman,sids in [('I',['01','02','03']),('II',['04']),('III',['05']),('IV',['06'])]:
    stem=f'George-Douglas-Annotated-Transcription-Volume-{roman}'
    pdf_path=WORK/f'render-{roman}'/f'{stem}.pdf'
    reader=PdfReader(pdf_path)
    volume_counts[roman]=len(reader.pages)
    assert len(list((WORK/f'render-{roman}').glob('page-*.png')))==len(reader.pages)
    texts=[p.extract_text() for p in reader.pages]
    clean=[re.sub(r'GEORGE DOUGLAS PAPERS\s*\|\s*ANNOTATED TRANSCRIPTION\s*\|\s*VOLUME [IV]+','',t) for t in texts]
    clean=[re.sub(r'Version 1\.1\s*•\s*\d+','',t) for t in clean]
    pdftext='\n'.join(clean)
    assert '\ufffd' not in pdftext
    normalized=norm(pdftext)
    fallback_normalized=None
    doc=Document(OUT/f'{stem}.docx')
    doctext='\n'.join(p.text for p in doc.paragraphs)
    paragraph_map={p.text.split('   ',1)[0]:p.text.split('   ',1)[1] for p in doc.paragraphs if re.match(r'^\d{2}-\d{3}-P\d+   ',p.text)}
    selected=[p for p in pages if p['source'] in sids]
    for p in selected:
        headings=[]
        for n,t in enumerate(texts,1):
            if re.search(r'Scan\s+'+f"{p['scan_page']:03}"+r'\s+'+re.escape(p['id'])+r'\b',t):
                headings.append(n)
        assert len(headings)==1,(p['id'],headings)
        page_map.append({'page_id':p['id'],'source_pdf':next(s['filename'] for s in sources if s['id']==p['source']),'source_scan_page':p['scan_page'],'edition_volume':roman,'edition_pdf':f'{stem}.pdf','entry_start_pdf_page':headings[0]})
        for block in p['paragraphs']:
            assert paragraph_map[block['id']]==block['text'],block['id']
            assert len(re.findall(re.escape(block['id'])+r'\b',pdftext))>=1,block['id']  # Witness notes repeat referenced IDs.
            if norm(block['text']) not in normalized:
                if fallback_normalized is None:
                    with pdfplumber.open(pdf_path) as alternate:
                        fallback='\n'.join(p.extract_text() or '' for p in alternate.pages)
                    fallback=re.sub(r'GEORGE DOUGLAS PAPERS\s*\|\s*ANNOTATED TRANSCRIPTION\s*\|\s*VOLUME [IV]+','',fallback)
                    fallback=re.sub(r'Version 1\.1\s*•\s*\d+','',fallback)
                    fallback_normalized=norm(fallback)
                if norm(block['text']) not in fallback_normalized:
                    pdf_issues.append(block['id'])
        if p['editorial_notes']:
            assert p['editorial_notes'] in doctext,p['id']
            if norm(p['editorial_notes']) not in normalized:
                pdf_issues.append(p['id']+' note')
    shutil.copy2(pdf_path,OUT/f'{stem}.pdf')
    (WORK/f'render-{roman}'/'extracted.txt').write_text('\n'.join(texts))
assert not pdf_issues,pdf_issues
with (SUP/'PDF-page-map.tsv').open('w',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=list(page_map[0]),delimiter='\t')
    writer.writeheader();writer.writerows(page_map)

plain=(SUP/'Complete-transcript.txt').read_text()
tree=html.fromstring((SUP/'Complete-transcript.html').read_text())
ids=tree.xpath('//*[@id]/@id')
assert len(ids)==len(set(ids))
assert len(tree.xpath('//article'))==387
for p in pages:
    for b in p['paragraphs']:
        assert b['text'] in plain,b['id']
        element=tree.get_element_by_id(b['id'])
        assert element.text_content()==b['id']+b['text'],b['id']
register=list(csv.DictReader((SUP/'Unresolved-readings.tsv').open(),delimiter='\t'))
assert len(register)==1175
assert len({r['item'] for r in register})==len(register)
regtree=html.fromstring((SUP/'Unresolved-readings.html').read_text())
assert len(regtree.xpath('//tbody/tr'))==len(register)
for link in regtree.xpath('//tbody//a/@href'):
    file_name,anchor=link.split('#')
    assert file_name=='Complete-transcript.html' and anchor in ids,link
assert len(list(csv.DictReader((SUP/'Page-inventory.tsv').open(),delimiter='\t')))==387
for t in (tree,regtree):
    assert not t.xpath('//script[@src] | //img[@src] | //link[@href]')
cases=[]
for name,t,selector in [('Complete-transcript.html',tree,'article'),('Unresolved-readings.html',regtree,'tbody tr')]:
    elements=t.xpath('//article') if selector=='article' else t.xpath('//tbody/tr')
    cases.append({'name':name,'script':'\n'.join(t.xpath('//script/text()')),'selector':selector,'texts':[e.text_content() for e in elements]})
(WORK/'html-filter-cases.json').write_text(json.dumps(cases,ensure_ascii=False))

readme=(OUT/'README.txt').read_text().replace('Supporting-files/PDF-page-map.tsv locates each source-page entry in the four\nedition PDFs, including entries that continue across more than one page.','Supporting-files/PDF-page-map.tsv gives the starting edition PDF page for\neach source-page entry. Some entries continue onto the following PDF page.')
(OUT/'README.txt').write_text(readme)
review_counts=dict(collections.Counter(p['review'] for p in pages))
audit={'source_scan_pages':len(pages),'transcript_blocks':len(block_ids),'unresolved_items':len(register),'edition_pdf_pages':volume_counts,'source_checksums_verified':True,'docx_text_matches_master':True,'pdf_transcript_text_matches_master':True,'html_and_text_match_master':True,'html_anchors_verified':True,'review_levels':review_counts,'visual_review':'Pending final visual inspection; see visual-review.json when complete.'}
(WORK/'final-audit.json').write_text(json.dumps(audit,indent=2)+'\n')
print(json.dumps(audit,indent=2))
