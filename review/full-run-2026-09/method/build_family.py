from pathlib import Path
import re,json,collections
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
BASE=Path('/Users/volneydouglas/Documents/FamilyHistory')
SRC=BASE/'book/george-douglas-independent';OUT=BASE/'output/george-douglas-independent'
files=[('01-finding-aid.md',6,'Collection finding aid'),('02-sermon.md',4,'Sermon notes'),('03-endorsement.md',5,'Sylvana Church endorsement'),('04-folder1.md',96,'Folder 1'),('05-folder2.md',158,'Folder 2'),('06-folder3.md',118,'Folder 3')]
entries={};ledger=[]
for fn,total,desc in files:
 source=fn[:2];text=(SRC/fn).read_text();blocks=re.split(r'^## Pages (\d+)-(\d+)\s*\n',text,flags=re.M);items=[];coverage=collections.Counter()
 for i in range(1,len(blocks),3):
  a,b=int(blocks[i]),int(blocks[i+1]);body=blocks[i+2].strip();m=re.match(r'### ([^\n]+)\n+(.*)',body,re.S);assert m,(fn,a)
  title,body=m.groups();assert len(title)<250,(fn,a,title);coverage.update(range(a,b+1));item=dict(source=source,source_name=desc,start=a,end=b,title=title,body=body,anchor=f's{source}p{a}')
  items.append(item);ledger.append({k:v for k,v in item.items() if k!='body'})
 assert set(coverage)==set(range(1,total+1)),(fn,coverage)
 assert all(c==1 for c in coverage.values()),fn
 entries[source]=sorted(items,key=lambda x:x['start'])
assert sum(x[1] for x in files)==387

def bookmark(p,name):
 start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bookmark.n));start.set(qn('w:name'),name)
 end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark.n));bookmark.n+=1;p._p.insert(1 if p._p.pPr is not None else 0,start);p._p.append(end)
bookmark.n=1

def addlink(p,text,anchor):
 h=OxmlElement('w:hyperlink');h.set(qn('w:anchor'),anchor);r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');rp.append(color);r.append(rp);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def addfield(p,code):
 r=p.add_run();x=OxmlElement('w:fldSimple');x.set(qn('w:instr'),code);r._r.addnext(x)

def setup(num,title):
 d=Document();sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.85);sec.header_distance=sec.footer_distance=Inches(.33)
 s=d.styles['Normal'];s.font.name='Georgia';s.font.size=Pt(11.5);s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.line_spacing=1.13;s.paragraph_format.space_after=Pt(7);s.paragraph_format.widow_control=True
 for name,size in [('Title',32),('Subtitle',17),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
  s=d.styles[name];s.font.name='Georgia';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.paragraph_format.keep_with_next=True;s.paragraph_format.space_before=Pt(14);s.paragraph_format.space_after=Pt(7)
 for name in ['Caption','Header','Footer']:
  d.styles[name].font.name='Calibri';d.styles[name].font.color.rgb=RGBColor(0,0,0)
 d.styles['Caption'].font.size=Pt(9);d.styles['Caption'].font.italic=False
 d.styles['Caption'].paragraph_format.keep_with_next=True
 sec.different_first_page_header_footer=True
 h=sec.header.paragraphs[0];h.text=f'GEORGE DOUGLAS PAPERS  ·  VOLUME {num}';h.style=d.styles['Header'];h.runs[0].font.size=Pt(8)
 f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('Independent reading edition  •  ');addfield(f,'PAGE');f.runs[0].font.size=Pt(8)
 p=d.add_paragraph();p.paragraph_format.space_before=Inches(1.3);r=p.add_run('THE GEORGE DOUGLAS PAPERS');r.font.name='Calibri';r.font.size=Pt(11);r.bold=True
 d.add_paragraph('A family reading\nedition','Title');d.add_paragraph(f'Volume {num}\n{title}','Subtitle')
 p=d.add_paragraph('A revised reading in plain English');p.paragraph_format.space_before=Pt(20)
 d.add_paragraph('Prepared for Volney Douglas\nRevision 1.1 September 2026')
 p=d.add_paragraph('Read from the scanned originals. Uncertain words and damaged passages are marked in brackets.');p.paragraph_format.space_before=Pt(34);p.runs[0].font.size=Pt(10)
 d.add_page_break();d.core_properties.title=f'George Douglas Papers — Volume {num}: {title}';d.core_properties.subject='Independent plain-English family reading edition';d.core_properties.author='Prepared for Volney Douglas';d.core_properties.keywords='George Douglas; family history; independent reading; manuscripts'
 for st in d.styles:
  if st._element.rPr is not None:
   fonts=st._element.rPr.find(qn('w:rFonts'))
   if fonts is not None:
    for key in list(fonts.attrib):
     if 'theme' in key.lower():del fonts.attrib[key]
  if st._element.pPr is not None:
   for border in list(st._element.pPr.findall(qn('w:pBdr'))):st._element.pPr.remove(border)
 return d

def table(d,block):
 rows=[[c.strip() for c in l.strip('|').split('|')] for l in block.splitlines() if l.strip()]
 rows=[r for r in rows if not all(re.fullmatch(r'[:\- ]+',c) for c in r)]
 t=d.add_table(rows=0,cols=len(rows[0]));t.autofit=False;t.columns[0].width=Inches(4.8);t.columns[1].width=Inches(1.8)
 pr=t._tbl.tblPr;bord=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  x=OxmlElement(f'w:{edge}');x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');bord.append(x)
 pr.append(bord)
 for idx,vals in enumerate(rows):
  row=t.add_row();trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement('w:cantSplit'))
  if idx==0:trPr.append(OxmlElement('w:tblHeader'))
  for j,val in enumerate(vals):
   c=row.cells[j];c.text=val
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
    for r in p.runs:r.font.size=Pt(10.5);r.bold=idx==0
   if idx==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EFEDE8');c._tc.get_or_add_tcPr().append(shade)
 d.add_paragraph().paragraph_format.space_after=Pt(0)

def markdown(d,content):
 for block in re.split(r'\n\s*\n',content.strip()):
  if block.startswith('# '):d.add_heading(block[2:],level=1);continue
  if block.startswith('## '):d.add_heading(block[3:],level=2);continue
  if block.startswith('### '):d.add_heading(block[4:],level=3);continue
  if block.startswith('|'):table(d,block);continue
  if block.startswith('- '):
   for line in block.splitlines():
    p=d.add_paragraph(line.removeprefix('- '),'List Bullet');p.paragraph_format.space_after=Pt(4)
   continue
  p=d.add_paragraph();p.add_run(block)
  if re.fullmatch(r'\[(?:Scan page|Opening of scan|Remaining text of scan)[^\]]*\]',block):p.paragraph_format.keep_with_next=True
  if block.startswith('[') and block.endswith(']'):
   for r in p.runs:r.italic=True;r.font.size=Pt(10.5)
  if '\n' in block and len(block.splitlines())>=4:
   p.paragraph_format.keep_together=True

def contents(d,items):
 d.add_heading('Contents',1)
 d.add_paragraph('Entries are arranged by source scan page. Select an entry to open it. The page numbers below refer to the original scan, not this edition.')
 for e in items:
  pg=str(e['start']) if e['start']==e['end'] else f"{e['start']}–{e['end']}"
  p=d.add_paragraph();p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.05
  addlink(p,f"{e['source_name']} · {pg}   |   {e['title']}",e['anchor'])
  p.style=d.styles['Normal']
  pp=p._p.get_or_add_pPr();size=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');size.append(sz);pp.append(size)
  for r in p._p.xpath('./w:hyperlink/w:r'):
   rp=r.find(qn('w:rPr'));sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');rp.append(sz)
 d.add_page_break()

def sourceentry(d,e):
 pg=str(e['start']) if e['start']==e['end'] else f"{e['start']}–{e['end']}"
 p=d.add_paragraph(f"SOURCE {e['source']} · {e['source_name'].upper()} · SCAN {'PAGE' if e['start']==e['end'] else 'PAGES'} {pg}",'Caption');p.paragraph_format.space_before=Pt(15)
 p=d.add_heading(e['title'],2);bookmark(p,e['anchor']);markdown(d,e['body'])

volumes=[('I','Reading guide and church papers',['01','02','03'],'01-Guide-and-Church-Papers'),('II','Folder 1: correspondence and loan papers',['04'],'02-Folder-1'),('III','Folder 2: correspondence and legal papers',['05'],'03-Folder-2'),('IV','Folder 3: correspondence and legal papers',['06'],'04-Folder-3')]
for num,title,sources,stem in volumes:
 d=setup(num,title)
 if num=='I':
  markdown(d,(SRC/'00-reading-guide.md').read_text());d.add_page_break();d.add_heading('Source files',1)
  source_names=sorted(p.name for p in Path('/Users/volneydouglas/Library/Mobile Documents/com~apple~CloudDocs/George Douglas').glob('ACU*.pdf'))
  for name,(_,n,_) in zip(source_names,files):d.add_paragraph(f'{name}\n{n} scan pages')
  d.add_paragraph('Total: 387 primary-source scan pages. The existing transcription PDF is excluded.');d.add_page_break()
 else:
  d.add_heading('About this volume',1)
  d.add_paragraph(f'This volume is a readable modern-English version of {title.lower()}. It accounts for every page in the supplied scan, including covers, notes, fragments, and duplicate documents.')
  d.add_paragraph('Square brackets mark uncertain words, lost passages, and editorial explanations. Page references in the entries count the cover of the source PDF as page 1. The source images, rather than this modernization, remain the basis for settling differences between readings.')
  d.add_paragraph('The first reading was prepared separately. Revision 1.1 incorporates selected source-image checks informed by comparison with the annotated edition and another agent’s contextual review. It is AI-assisted and has not been independently human verified. See Volume I for the method and limits.')
  d.add_paragraph('Use the linked contents, PDF bookmarks, or Word Navigation pane to move between documents. The printed footer numbers belong to this edition; the source-page labels lead back to the scans.');d.add_page_break()
 items=[e for s in sources for e in entries[s]];contents(d,items)
 for e in items:sourceentry(d,e)
 path=OUT/f'{stem}.docx';d.save(path);print(path.name,len(items),'entries')
(OUT/'Source-ledger.json').write_text(json.dumps(dict(source_page_count=387,entries=sorted(ledger,key=lambda e:(e['source'],e['start']))),indent=2,ensure_ascii=False))
(OUT/'Read-me-first.txt').write_text('GEORGE DOUGLAS PAPERS — INDEPENDENT FAMILY READING EDITION\n\nStart with 01-Guide-and-Church-Papers.pdf.\n\nFour PDFs are ready to read and share; four matching Word documents are editable. All 387 pages of the six supplied primary PDFs are accounted for. Illegible passages remain marked. The existing transcription PDF is excluded. The archival finding aid describes additional folders not supplied among the local scans.\n\nRead the independence note and editorial conventions in Volume I before comparing readings. Scan-page references accompany every entry. Source-ledger.json provides a machine-readable coverage index.\n\nPrepared for Volney Douglas, September 2026. Original source PDFs were not altered.\n')
