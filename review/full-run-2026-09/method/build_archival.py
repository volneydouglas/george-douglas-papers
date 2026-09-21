from pathlib import Path
import json,re,csv,html,hashlib,shutil,collections
from datetime import datetime,timezone
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE=Path('/Users/volneydouglas/Documents/FamilyHistory')
BOOK=BASE/'book/george-douglas-archival'
OUT=BASE/'output/george-douglas-archival';OUT.mkdir(parents=True,exist_ok=True)
SUP=OUT/'Supporting-files';SUP.mkdir(exist_ok=True)
SOURCES=json.loads((BOOK/'sources.json').read_text())
PAGES=[json.loads(p.read_text()) for p in sorted((BOOK/'pages').glob('*.json'))]
assert len(PAGES)==387
assert [d['id'] for d in PAGES]==[f"{s['id']}-{n:03}" for s in SOURCES for n in range(1,s['pages']+1)]
for d in PAGES:
 assert d['text'].strip() and d['text'].count('[[')==d['text'].count(']]'),d['id']
 d['paragraphs']=[{'id':f"{d['id']}-P{n:02}",'text':p} for n,p in enumerate(d['text'].split('\n\n'),1)]
 d['review']=d.get('review','initial_image_reading')
SOURCEDICT={s['id']:s for s in SOURCES}
VOLUMES=[('I','Guide and finding aid with sermon and church papers',['01','02','03']),('II','Folder 1 Correspondence and loan papers',['04']),('III','Folder 2 Correspondence and legal papers',['05']),('IV','Folder 3 Correspondence and legal papers',['06'])]
GUIDE=[
('Purpose and scope', 'This edition offers a direct, annotated transcription of the six supplied PDFs of the George Douglas Papers, 1830–1939, Center for Restoration Studies MS 5, Abilene Christian University Special Collections and Archives. It accounts for all 387 scan pages, including repository covers, later identifying slips, envelopes, duplicate images and fragments. Some damaged passages remain only partly transcribed. Folders 4–7 described in the finding aid were not supplied as manuscript scans and are outside this edition.'),
('Preparation and review','Prepared for Volney Douglas with AI assistance in September 2026, from visual reading of the supplied scans. Each scan received an initial reading. Version 1.1 adds local OCR processing of all 193 pages containing the original 1,200 uncertainties, and applies selected image-checked corrections. OCR completion is not a complete new visual check; the item ledger records which passages were individually adjudicated. The page inventory identifies the review level. This edition has not received an independent human transcription check or repository approval. It is offered for scholarly review and correction, and does not replace the source images.'),
('Independence','The first edition was prepared separately from the plain-English family edition. For version 1.1, the other agent’s contextual review and the family edition were consulted to guide image checks; this revision is not a blinded comparison. Part of a pre-existing 25-page transcription was viewed during the initial inventory, before the request for independence; it has not been used for this transcription. The work is therefore not a blind experiment.'),
('Wording and layout','Original wording, spelling, grammar, abbreviations and recoverable punctuation are retained. No silent expansion or modernization is intended. Handwritten capitalization, punctuation and word division sometimes require judgment. Superscripts are lowered to the baseline. Ordinary line wrapping is reflowed; verse lines, addresses and separate text blocks are distinguished. Underlining and decorative features are not systematically encoded. This is a direct reading transcript, not a facsimile or a fully diplomatic edition.'),
('Uncertainty notation','Double square brackets contain editorial intervention. [[illegible]] means writing could not be read; an unspecified gap does not imply a word count. Adjacent unreadable spans may be combined. A stated extent, such as [[illegible: approximately 3 words]], is an estimate. [[word?]] gives a tentative reading, including tentative phrases. [[lost: torn edge]] records writing missing at a damaged edge; [[obscured: fold]] identifies an obstruction. Original single brackets remain as written. No missing wording is supplied solely to produce a fluent sentence.'),
('Corrections and separate blocks','[[deleted: text]] records readable canceled text; [[deleted: illegible]] records an unread cancellation. [[inserted: text]] records a significant insertion in its intended position. [[margin]] and [[/margin]] enclose marginal writing. [[postmark]] and [[postage]] identify postal text. [[blank]], [[seal]] and [[sketch]] identify useful non-textual features. A [[separate block: description]] marker introduces an independently positioned passage. Handwriting or cancellation may itself be uncertain.'),
('Source order and later descriptions','Entries follow the PDFs’ scan order. Where the leaves appear out of sequence, the editorial note suggests a reading order without moving the entry. Later identifying slips are transcribed separately; their names, dates and attributions are not silently substituted for manuscript readings. Duplicate scans retain separate entries. Apparent contradictions are recorded rather than reconciled.'),
('Documentary statements','The letters preserve their writers’ assertions, beliefs, rumors and opinions. Transcription does not establish those statements as historical fact. Historical racial language and other offensive wording remain where legible. The original scans are the evidence against which this typescript should be checked.'),
('References and corrections','A page ID such as 05-025 means source 05, PDF scan page 25, counting the cover as page 1. A block ID such as 05-025-P02 adds the second transcript paragraph or block. These are editorial locators; they are not manuscript page numbers. Suggested citation: George Douglas Papers annotated transcription, version 1.1, 05-025-P02; corresponding source 05, scan page 25. For a correction, give that locator, the existing wording, a proposed reading, and the image or other evidence supporting it.'),
('Files for reuse','The package includes four searchable PDFs and four editable Word files, a complete UTF-8 transcript, structured JSON, an offline HTML reading copy, a page inventory, and an unresolved-readings register in TSV and HTML. The register retains the original uncertainty IDs; resolved IDs are retired and new IDs use higher numbers. The review ledger preserves all 1,200 baseline items. The register includes tentative readings, unread text, loss and obstruction; structural markers and readable insertions or deletions are not unresolved items. Original source PDFs are identified by filename and SHA-256 checksum but are not duplicated in this package.'),
('Repository credit','Original documents and scans: George Douglas Papers, 1830–1939, Center for Restoration Studies MS 5, Abilene Christian University Special Collections and Archives, Brown Library, Abilene Christian University, Abilene, Texas. Collection address printed in the supplied files: https://digitalcommons.acu.edu/george_douglas. This contribution acknowledges the repository’s work in preserving and scanning the collection. No repository endorsement or transfer of copyright is asserted.'),
('Transcription guidance','The general approach of retaining original spelling and abbreviations and marking uncertainty is informed by the Smithsonian Transcription Center’s How to Transcribe guidance, consulted September 2026: https://transcription.si.edu/instructions/transcribe. This edition defines its own notation and is not a Smithsonian project.')]
(BOOK/'EDITORIAL-GUIDE.md').write_text('# George Douglas Papers annotated transcription\n\nVersion 1.1 • September 2026\n\n'+'\n\n'.join('## '+h+'\n\n'+t for h,t in GUIDE)+'\n')

# Deterministic, stable source and paragraph locators also appear in reusable text.
jsonobj={'title':'George Douglas Papers annotated transcription','version':'1.1','date':'2026-09','scope_scan_pages':387,'editorial_guide':[{'heading':h,'text':t} for h,t in GUIDE],'sources':SOURCES,'pages':PAGES}
(SUP/'Transcripts.json').write_text(json.dumps(jsonobj,ensure_ascii=False,indent=2)+'\n')
(SUP/'Source-manifest.json').write_text(json.dumps(SOURCES,ensure_ascii=False,indent=2)+'\n')
shutil.copy2(BOOK/'EDITORIAL-GUIDE.md',SUP/'Editorial-guide.md')
plain=['GEORGE DOUGLAS PAPERS ANNOTATED TRANSCRIPTION','Version 1.1 • September 2026','']
for h,t in GUIDE:plain.extend([h.upper(),t,''])
for d in PAGES:
 plain.extend(['='*68,d['id']+' | '+d['kind'],'Source: '+SOURCEDICT[d['source']]['filename']+f" | PDF scan page {d['scan_page']}",''])
 for p in d['paragraphs']:plain.extend([p['id'],p['text'],''])
 if d['editorial_notes']:plain.extend(['EDITORIAL NOTE: '+d['editorial_notes'],''])
(SUP/'Complete-transcript.txt').write_text('\n'.join(plain)+'\n')
index=json.loads((BOOK/'uncertainty-index.json').read_text())
index_by_block=collections.defaultdict(list)
for row in index:index_by_block[row['paragraph_id']].append(row)
register=[]
for d in PAGES:
 seq=0
 for p in d['paragraphs']:
  for m in re.finditer(r'\[\[(.*?)\]\]',p['text'],re.S):
   ann=m.group(1)
   if not ('?' in ann or 'illegible' in ann or ann.startswith(('lost:','obscured:'))):continue
   seq+=1
   indexed=index_by_block[p['id']].pop(0)
   assert indexed['annotation']==m.group(0),(p['id'],indexed,m.group(0))
   category='tentative reading' if '?' in ann else 'loss' if ann.startswith('lost:') else 'obstruction' if ann.startswith('obscured:') else 'unread writing'
   context=p['text'][max(0,m.start()-65):min(len(p['text']),m.end()+65)].replace('\n',' / ')
   register.append({'item':indexed['item'],'page_id':d['id'],'paragraph_id':p['id'],'source_pdf':SOURCEDICT[d['source']]['filename'],'scan_page':d['scan_page'],'category':category,'annotation':m.group(0),'context':context,'proposed_correction':indexed.get('proposed_correction',''),'reviewer_and_date':indexed.get('reviewer_and_date',''),'evidence':indexed.get('evidence','')})
with (SUP/'Unresolved-readings.tsv').open('w',newline='') as f:
 w=csv.DictWriter(f,fieldnames=list(register[0]),delimiter='\t');w.writeheader();w.writerows(register)
with (SUP/'Page-inventory.tsv').open('w',newline='') as f:
 w=csv.writer(f,delimiter='\t');w.writerow(['page_id','source_pdf','scan_page','kind','transcript_blocks','unresolved_items','review_level','editorial_note'])
 for d in PAGES:w.writerow([d['id'],SOURCEDICT[d['source']]['filename'],d['scan_page'],d['kind'],len(d['paragraphs']),sum(x['page_id']==d['id'] for x in register),d['review'],d['editorial_notes']])
(SUP/'Review-log.jsonl').write_text((BOOK/'review-log.jsonl').read_text())

CSS='''body{font:18px/1.55 Georgia,serif;color:#202020;max-width:900px;margin:3rem auto;padding:0 1.5rem}h1,h2,h3,nav,label,.meta{font-family:Arial,sans-serif}h1{font-size:2rem;line-height:1.15}h2{font-size:1.25rem;margin-top:2.5rem}.meta{font-size:.8rem;color:#555}.note{font:14px/1.5 Arial,sans-serif;background:#f3f3f3;padding:1rem}a{color:#234b68}article{padding:1rem 0;border-bottom:1px solid #ddd}input,select{font:16px Arial;padding:.5rem;width:100%;box-sizing:border-box;margin:.5rem 0}nav{position:sticky;top:0;background:white;border-bottom:1px solid #bbb;padding:.5rem 0;z-index:1}.uncertain{background:#fff4d6}.locator{font:12px Arial;color:#555;margin-right:.6rem}.block{white-space:pre-wrap;margin:1rem 0}table{width:100%;border-collapse:collapse;font:14px/1.45 Arial}th,td{border-bottom:1px solid #ddd;padding:.5rem;text-align:left;vertical-align:top}th{background:#eee}button{font:16px Arial;padding:.4rem}article[hidden],tr[hidden]{display:none}@media print{nav{position:static}input,button,select{display:none}body{max-width:none;font-size:11pt}.note{background:none}h2{break-after:avoid}}'''
def E(s):return html.escape(str(s))
def markup(t):return re.sub(r'(\[\[.*?\]\])',r'<span class="uncertain">\1</span>',E(t),flags=re.S)
reader=['<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>George Douglas Papers annotated transcription</title><style>'+CSS+'</style><body><h1>George Douglas Papers annotated transcription</h1><p>Version 1.1 · September 2026 · 387 source scan pages</p><p>AI-assisted contribution for repository review. Unreadable and tentative passages remain marked. <a href="Unresolved-readings.html">Open the unresolved-readings register</a>.</p><details><summary>Editorial guide</summary>']
reader.extend('<h2>'+E(h)+'</h2><p>'+E(t)+'</p>' for h,t in GUIDE);reader.append('</details><nav><label for="q">Find wording or a source page ID</label><input id="q" type="search" placeholder="For example 06-063 or a surname"><span class="meta" id="count"></span></nav>')
for d in PAGES:
 reader.append(f'<article id="{d["id"]}"><h2>{d["id"]} · {E(d["kind"])}</h2><p class="meta">{E(SOURCEDICT[d["source"]]["filename"])} · Scan page {d["scan_page"]}</p>')
 for p in d['paragraphs']:reader.append(f'<div class="block" id="{p["id"]}"><span class="locator">{p["id"]}</span>{markup(p["text"])}</div>')
 if d['editorial_notes']:reader.append('<p class="note"><strong>Editorial note</strong> '+E(d['editorial_notes'])+'</p>')
 reader.append('</article>')
reader.append('''<script>const articles=[...document.querySelectorAll('article')],q=document.getElementById('q');function filter(){const s=q.value.toLocaleLowerCase();let n=0;articles.forEach(a=>{a.hidden=!a.textContent.toLocaleLowerCase().includes(s);if(!a.hidden)n++});document.getElementById('count').textContent=n+' of 387 source pages shown';}q.addEventListener('input',filter);filter();</script></body></html>''')
(SUP/'Complete-transcript.html').write_text('\n'.join(reader))
reg=['<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>George Douglas Papers unresolved readings</title><style>'+CSS+'body{max-width:1150px}</style><body><h1>Unresolved readings</h1><p>George Douglas Papers annotated transcription · Version 1.1</p><p>Each item links to its transcript paragraph. The TSV copy retains contextual proposals and review evidence. Original uncertainty IDs remain stable; retired IDs appear in the revision ledger. This register records uncertainty in transcription, not the truth of a writer’s statements.</p><nav><label for="q">Filter by wording or locator</label><input id="q" type="search"><span class="meta" id="count"></span></nav><table><thead><tr><th>Locator</th><th>Type</th><th>Reading and context</th></tr></thead><tbody>']
for r in register:reg.append('<tr><td><a href="Complete-transcript.html#'+r['paragraph_id']+'">'+r['item']+'</a><br>'+r['paragraph_id']+'</td><td>'+r['category']+'</td><td><strong>'+E(r['annotation'])+'</strong><br>'+E(r['context'])+'<br><em>Proposal:</em> '+E(r['proposed_correction'])+'<br><em>Review:</em> '+E(r['evidence'])+'</td></tr>')
reg.append('''</tbody></table><script>const rows=[...document.querySelectorAll('tbody tr')],q=document.getElementById('q');function filter(){let n=0;const s=q.value.toLocaleLowerCase();rows.forEach(r=>{r.hidden=!r.textContent.toLocaleLowerCase().includes(s);if(!r.hidden)n++});document.getElementById('count').textContent=n+' of '+rows.length+' items shown';}q.addEventListener('input',filter);filter();</script></body></html>''')
(SUP/'Unresolved-readings.html').write_text('\n'.join(reg))

# A simple scholarly layout, with page IDs as Word navigation headings.
def font(style,name,size,bold=False,color='000000'):
 style.font.name=name;style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor.from_string(color)
 style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
def field(p,instruction):
 r=p.add_run();b=OxmlElement('w:fldChar');b.set(qn('w:fldCharType'),'begin');r._r.append(b)
 t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=' '+instruction+' ';r._r.append(t)
 s=OxmlElement('w:fldChar');s.set(qn('w:fldCharType'),'end');r._r.append(s)
def addtext(p,text):
 for part in re.split(r'(\[\[.*?\]\])',text,flags=re.S):
  r=p.add_run(part)
  if part.startswith('[['):r.italic=True

def base_doc(roman):
 doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.78);sec.bottom_margin=Inches(.75);sec.left_margin=Inches(.85);sec.right_margin=Inches(.85);sec.header_distance=Inches(.3);sec.footer_distance=Inches(.3)
 st=doc.styles
 for border in st.element.xpath('.//w:pBdr'):
  border.getparent().remove(border)
 font(st['Normal'],'Liberation Serif',11.5);st['Normal'].paragraph_format.line_spacing=1.13;st['Normal'].paragraph_format.space_after=Pt(6)
 for name,size in [('Title',28),('Subtitle',16),('Heading 1',18),('Heading 2',13)]:
  font(st[name],'Liberation Sans',size,bold=name!='Subtitle');st[name].paragraph_format.space_before=Pt(15 if name=='Heading 2' else 12);st[name].paragraph_format.space_after=Pt(7)
 for name,size in [('Metadata',9),('Editorial Note',9.5),('Guide Body',10.5)]:
  style=st.add_style(name,WD_STYLE_TYPE.PARAGRAPH);style.base_style=st['Normal'];font(style,'Liberation Sans',size,color='333333');style.paragraph_format.line_spacing=1.12
 st['Metadata'].paragraph_format.space_after=Pt(7);st['Metadata'].paragraph_format.keep_with_next=True
 st['Editorial Note'].paragraph_format.space_before=Pt(4);st['Editorial Note'].paragraph_format.space_after=Pt(10)
 st['Normal'].paragraph_format.widow_control=True
 h=sec.header.paragraphs[0];h.text=f'GEORGE DOUGLAS PAPERS   |   ANNOTATED TRANSCRIPTION   |   VOLUME {roman}';font(st['Header'],'Liberation Sans',8,color='555555')
 f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('Version 1.1  •  ');field(f,'PAGE');font(st['Footer'],'Liberation Sans',8,color='555555')
 doc.core_properties.title=f'George Douglas Papers annotated transcription Volume {roman}';doc.core_properties.author='Prepared for Volney Douglas with AI assistance';doc.core_properties.subject='Contribution for ACU Special Collections and Archives review';doc.core_properties.keywords='George Douglas; ACU; MS 5; transcription; genealogy';doc.core_properties.comments='Source scan pages and uncertainties are identified in the text.'
 doc.core_properties.created=datetime.now(timezone.utc);doc.core_properties.modified=doc.core_properties.created;doc.core_properties.last_modified_by='AI-assisted transcription for Volney Douglas';doc.core_properties.language='en-US'
 return doc
for roman,subtitle,sids in VOLUMES:
 doc=base_doc(roman)
 doc.add_paragraph('George Douglas Papers\nannotated transcription',style='Title')
 doc.add_paragraph('Volume '+roman,style='Subtitle')
 doc.add_paragraph(subtitle,style='Subtitle')
 doc.add_paragraph('Version 1.1\nSeptember 2026',style='Guide Body')
 doc.add_paragraph('Prepared for Volney Douglas\nFor contribution to Abilene Christian University\nSpecial Collections and Archives',style='Guide Body')
 doc.add_paragraph('Direct transcription with marked uncertainties and editorial notes. Original scans supplied by ACU; AI-assisted typescript offered for repository review.',style='Guide Body')
 count=sum(SOURCEDICT[x]['pages'] for x in sids)
 doc.add_paragraph(f'This volume covers {count} source scan pages. The complete set covers six PDFs and 387 scan pages. Damaged passages remain partial or tentative where marked.',style='Guide Body')
 doc.add_heading('Volumes in this edition',level=2)
 for rn,sub,ids in VOLUMES:doc.add_paragraph(f'{rn}   {sub}',style='Guide Body')
 doc.add_page_break()
 if roman=='I':
  doc.add_heading('Editorial guide',level=1)
  for h,t in GUIDE:
   doc.add_heading(h,level=2);doc.add_paragraph(t,style='Guide Body')
  doc.add_heading('Source inventory',level=1)
  for s in SOURCES:
   doc.add_heading('Source '+s['id'],level=2);doc.add_paragraph(s['filename'],style='Guide Body');doc.add_paragraph(f"PDF scan pages: {s['pages']}\nSHA-256: {s['sha256']}",style='Metadata')
 else:
  doc.add_heading('Reading this volume',level=1)
  for index in [1,3,4,5,6,8,10]:
   h,t=GUIDE[index];doc.add_heading(h,level=2);doc.add_paragraph(t,style='Guide Body')
 for sid in sids:
  doc.add_page_break();doc.add_heading('Source '+sid,level=1)
  s=SOURCEDICT[sid];doc.add_paragraph(s['filename'],style='Guide Body');doc.add_paragraph(f"{s['pages']} PDF scan pages including the repository cover.\nSHA-256: {s['sha256']}",style='Metadata')
  for d in (x for x in PAGES if x['source']==sid):
   doc.add_heading('Scan '+str(d['scan_page']).zfill(3)+'   '+d['id'],level=2)
   doc.add_paragraph(d['kind'].capitalize()+'  |  '+f"Source {sid}, PDF page {d['scan_page']}",style='Metadata')
   for block in d['paragraphs']:
    p=doc.add_paragraph();p.paragraph_format.keep_together=False
    r=p.add_run(block['id']+'   ');r.font.name='Liberation Sans';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('555555')
    addtext(p,block['text'])
   if d['editorial_notes']:
    p=doc.add_paragraph(style='Editorial Note');p.add_run('Editorial note. ').bold=True;p.add_run(d['editorial_notes'])
 name=f'George-Douglas-Annotated-Transcription-Volume-{roman}.docx';doc.save(OUT/name)
 print(name)

(OUT/'README.txt').write_text('''GEORGE DOUGLAS PAPERS ANNOTATED TRANSCRIPTION
Version 1.1 • September 2026

Prepared for Volney Douglas as an AI-assisted contribution for ACU review.

Start with Volume I for the editorial guide, source inventory and finding aid.
Volumes II–IV contain the supplied manuscript folders 1–3 in PDF scan order.
All 387 supplied scan pages have entries. This is not a claim that every word
is legible: uncertain and unread passages remain marked. This edition has not
received independent human or repository verification.

Each volume is supplied as a searchable PDF and an editable DOCX.
For offline searching across all volumes, open:
  Supporting-files/Complete-transcript.html
For doubts and proposed corrections, open:
  Supporting-files/Unresolved-readings.html
The matching TSV preserves proposed readings, reviewers and evidence.
Review-summary.txt explains the revision and points to its item ledger.
Supporting-files/PDF-page-map.tsv gives the starting edition PDF page for
each source-page entry. Some entries continue onto the following PDF page.

Reusable UTF-8 text and JSON, a page inventory, source checksums, editorial
guide and focused-review log are in Supporting-files. Every source page has a
stable ID (for example 05-025); each transcript block adds P01, P02, and so on.
The source PDFs are not duplicated in this package. Use their exact filenames
and SHA-256 checksums to identify the scans. The original images remain the
authoritative evidence. The finding aid describes additional folders 4–7 that
were not supplied as manuscript scans and are outside this edition.

Draft-cover-note.txt is a suggested message to accompany the contribution.
Nothing has been sent to the university.

manifest.txt lists the package files using relative paths. SHA256SUMS.txt
records checksums for every package file except the checksum file itself.
''')
(OUT/'Draft-cover-note.txt').write_text('''Dear Special Collections and Archives team,

Thank you for the time and care you put into scanning the George Douglas
Papers for our family. Being able to read these records means a great deal
to us, and I hope we can return something useful for future researchers.

I am sharing an annotated transcription of the six PDFs you supplied. It
includes the original wording where readable, marks doubtful or unreadable
passages, and provides source-page and paragraph references. The package
contains searchable PDFs, editable Word files, reusable text and JSON, and
a register where corrections can be recorded.

This is an AI-assisted contribution for your review, rather than a verified
repository edition. Some damaged handwriting remains unresolved. The editorial
guide explains the method, limitations and source files, and the scans remain
the authoritative evidence. Please feel free to evaluate the materials and
adapt them to your cataloging and access practices as appropriate.

Thank you again for helping preserve these documents and make them accessible.

With appreciation,
Volney Douglas
''')
print('Scans',len(PAGES),'Blocks',sum(len(d['paragraphs']) for d in PAGES),'Unresolved items',len(register))
